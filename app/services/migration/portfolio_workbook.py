"""Apply a :mod:`portfolio_profile` to a workbook -> normalized records.

PURE-ish and DB-free: this layer knows spreadsheets and units, nothing about
PaySpyre's tables. It resolves each profile column binding against the file's
real header row, coerces cells into canonical Python types (integer CENTS,
basis points, ``date``), and returns dataclasses plus a per-row issue list. A
malformed row never raises — it is reported.

FIDELITY: transaction allocations (fees / interest / principal) and the running
balances are carried through EXACTLY as recorded, including negative values
(returns, refunds, corrections). Nothing here re-derives an allocation; the
historical record is the historical record.

Money is read as the profile's ``money_unit`` and stored as integer cents.
Spreadsheet floats carry binary noise (``1138.0000000000002``), so conversion
rounds half-up on the cent — never truncates.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from typing import Any, Iterable, Optional, Sequence

from app.services.migration.portfolio_profile import (
    PortfolioProfile,
    SheetSpec,
    TransposedSheetSpec,
)


# ---------------------------------------------------------------------------
# Workbook access
# ---------------------------------------------------------------------------


class WorkbookSource:
    """Minimal read interface a workbook must offer (keeps the reader testable
    without an .xlsx on disk)."""

    def sheet_names(self) -> list[str]:  # pragma: no cover - interface
        raise NotImplementedError

    def rows(self, sheet: str) -> list[tuple]:  # pragma: no cover - interface
        raise NotImplementedError


class InMemoryWorkbook(WorkbookSource):
    """``{sheet_name: [row_tuple, ...]}`` — what the tests use."""

    def __init__(self, sheets: dict[str, Sequence[Sequence[Any]]]):
        self._sheets = {k: [tuple(r) for r in v] for k, v in sheets.items()}

    def sheet_names(self) -> list[str]:
        return list(self._sheets)

    def rows(self, sheet: str) -> list[tuple]:
        return list(self._sheets.get(sheet, []))


class ExcelWorkbook(WorkbookSource):
    """An .xlsx opened with openpyxl (values only, formulas already evaluated)."""

    def __init__(self, path_or_bytes: Any):
        import openpyxl  # local import: keeps the module importable without the dep

        src = (
            io.BytesIO(path_or_bytes)
            if isinstance(path_or_bytes, (bytes, bytearray))
            else path_or_bytes
        )
        self._wb = openpyxl.load_workbook(src, data_only=True, read_only=True)

    def sheet_names(self) -> list[str]:
        return list(self._wb.sheetnames)

    def rows(self, sheet: str) -> list[tuple]:
        if sheet not in self._wb.sheetnames:
            return []
        return [tuple(r) for r in self._wb[sheet].iter_rows(values_only=True)]

    def close(self) -> None:
        self._wb.close()


# ---------------------------------------------------------------------------
# Cell coercion
# ---------------------------------------------------------------------------

_CENT = Decimal("0.01")
_DATE_FORMATS = ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%m-%d-%Y")


def to_cents(value: Any, *, money_unit: str = "dollars") -> Optional[int]:
    """Money cell -> integer cents (half-up), or None when there is no number.

    Spreadsheet money arrives as a float with binary noise; quantizing through
    ``Decimal(str(v))`` reproduces the number the operator SEES in the cell.
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        s = value.strip().replace("$", "").replace(",", "")
        if s.startswith("(") and s.endswith(")"):
            s = "-" + s[1:-1]
        if not s:
            return None
        try:
            value = Decimal(s)
        except Exception:
            return None
    if isinstance(value, (int, float, Decimal)):
        d = Decimal(str(value))
        if money_unit == "cents":
            return int(d.to_integral_value(rounding=ROUND_HALF_UP))
        return int((d * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    return None


def to_bps(value: Any, *, rate_unit: str = "fraction") -> Optional[int]:
    """Rate cell -> annual basis points."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        s = value.strip().rstrip("%")
        if not s:
            return None
        try:
            value = Decimal(s)
        except Exception:
            return None
    if not isinstance(value, (int, float, Decimal)):
        return None
    d = Decimal(str(value))
    factor = {"fraction": 10_000, "percent": 100, "bps": 1}[rate_unit]
    return int((d * factor).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def to_date(value: Any) -> Optional[date]:
    """Date cell -> ``date``. Non-date sentinels ("Closed", "") become None."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        s = value.strip()
        for fmt in _DATE_FORMATS:
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
    return None


def to_int(value: Any) -> Optional[int]:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return int(Decimal(str(value)).to_integral_value(rounding=ROUND_HALF_UP))
    if isinstance(value, str):
        s = value.strip()
        if not s:
            return None
        try:
            return int(Decimal(s).to_integral_value(rounding=ROUND_HALF_UP))
        except Exception:
            return None
    return None


def to_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, str):
        s = value.strip()
        return s or None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


# ---------------------------------------------------------------------------
# Normalized records
# ---------------------------------------------------------------------------


@dataclass
class PortfolioAccount:
    """One loan as the source system states it. Money in integer cents."""

    row_number: int
    account_number: str
    vendor_code: Optional[str] = None
    provider_name: Optional[str] = None
    source_status: str = ""
    source_sub_status: str = ""
    borrower_name: Optional[str] = None
    co_borrower_name: Optional[str] = None
    days_past_due: int = 0
    sales_value_cents: Optional[int] = None
    insurance_cents: Optional[int] = None
    down_payment_cents: Optional[int] = None
    amount_financed_cents: Optional[int] = None
    term_months: Optional[int] = None
    annual_rate_bps: Optional[int] = None
    regular_payment_cents: Optional[int] = None
    payment_frequency: Optional[str] = None       # source label, verbatim
    cost_of_borrowing_cents: Optional[int] = None
    first_payment_date: Optional[date] = None
    final_payment_date: Optional[date] = None
    days_in_year: Optional[int] = None
    origination_date: Optional[date] = None
    origination_type: Optional[str] = None
    # Stated cumulative actuals + balances — the reconciliation targets.
    payment_amount_total_cents: Optional[int] = None
    fees_paid_cents: Optional[int] = None
    interest_paid_cents: Optional[int] = None
    principal_paid_cents: Optional[int] = None
    fees_balance_cents: Optional[int] = None
    interest_balance_cents: Optional[int] = None
    principal_balance_cents: Optional[int] = None
    total_owed_cents: Optional[int] = None
    next_due_date: Optional[date] = None
    last_transaction_date: Optional[date] = None
    last_transaction_type: Optional[str] = None
    close_date: Optional[date] = None
    close_type: Optional[str] = None
    nsf_return_count: int = 0
    deferment_count: int = 0


@dataclass
class PortfolioTransaction:
    """One posted event, carried through VERBATIM. Money in integer cents.

    Allocations keep the sign the source recorded (a return is negative). This
    module never re-derives them.
    """

    row_number: int
    account_number: str
    source_type: str
    date: Optional[date]
    transaction_number: Optional[str] = None
    vendor_code: Optional[str] = None
    provider_name: Optional[str] = None
    payment_cents: Optional[int] = None
    fees_charged_cents: Optional[int] = None
    fees_paid_cents: Optional[int] = None
    fees_balance_cents: Optional[int] = None
    accrued_interest_cents: Optional[int] = None
    interest_due_cents: Optional[int] = None
    interest_paid_cents: Optional[int] = None
    interest_balance_cents: Optional[int] = None
    principal_paid_cents: Optional[int] = None
    principal_balance_cents: Optional[int] = None
    total_owed_cents: Optional[int] = None
    comment: Optional[str] = None


@dataclass
class PortfolioVendor:
    code: Optional[str] = None
    name: Optional[str] = None
    address: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    start_date: Optional[date] = None
    providers: list[str] = field(default_factory=list)


@dataclass
class ReadIssue:
    sheet: str
    row_number: Optional[int]
    message: str

    def as_dict(self) -> dict:
        return {"sheet": self.sheet, "row": self.row_number, "message": self.message}


@dataclass
class WorkbookReadResult:
    profile_name: str
    accounts: list[PortfolioAccount] = field(default_factory=list)
    transactions: list[PortfolioTransaction] = field(default_factory=list)
    vendors: list[PortfolioVendor] = field(default_factory=list)
    issues: list[ReadIssue] = field(default_factory=list)
    skipped_rows: int = 0

    def as_report(self) -> dict:
        return {
            "profile": self.profile_name,
            "accounts": len(self.accounts),
            "transactions": len(self.transactions),
            "vendors": len(self.vendors),
            "providers": sum(len(v.providers) for v in self.vendors),
            "skipped_rows": self.skipped_rows,
            "issues": [i.as_dict() for i in self.issues],
        }


class ProfileMismatch(Exception):
    """The profile's required columns are not present in the file."""


# ---------------------------------------------------------------------------
# Column resolution
# ---------------------------------------------------------------------------

_MONEY_FIELDS = {
    "sales_value", "insurance", "down_payment", "amount_financed",
    "regular_payment", "cost_of_borrowing", "payment_amount_total", "fees_paid",
    "interest_paid", "principal_paid", "fees_balance", "interest_balance",
    "principal_balance", "total_owed", "payment", "fees_charged",
    "accrued_interest", "interest_due",
}
_DATE_FIELDS = {
    "first_payment_date", "final_payment_date", "origination_date",
    "next_due_date", "last_transaction_date", "close_date", "date", "start_date",
}
_INT_FIELDS = {"days_past_due", "term_months", "days_in_year", "nsf_return_count", "deferment_count"}


def _normalize_header(v: Any) -> str:
    return " ".join(str(v).split()).casefold() if v is not None else ""


def resolve_columns(header: Sequence[Any], spec: SheetSpec) -> dict[str, int]:
    """canonical field -> 0-based column index, using header text or a literal index."""
    by_text: dict[str, int] = {}
    for idx, cell in enumerate(header):
        key = _normalize_header(cell)
        if key and key not in by_text:
            by_text[key] = idx
    resolved: dict[str, int] = {}
    for canonical, binding in spec.columns.items():
        if isinstance(binding, int):
            resolved[canonical] = binding
        else:
            idx = by_text.get(_normalize_header(binding))
            if idx is not None:
                resolved[canonical] = idx
    missing = [c for c in spec.required_columns if c not in resolved]
    if missing:
        raise ProfileMismatch(
            f"sheet {spec.sheet!r}: required column(s) not found in header row "
            f"{spec.header_row}: {', '.join(missing)}"
        )
    return resolved


def _cell(row: Sequence[Any], idx: Optional[int]) -> Any:
    if idx is None or idx < 0 or idx >= len(row):
        return None
    return row[idx]


def _coerce(field_name: str, raw: Any, profile: PortfolioProfile) -> Any:
    if field_name in _MONEY_FIELDS:
        return to_cents(raw, money_unit=profile.money_unit)
    if field_name == "annual_rate":
        return to_bps(raw, rate_unit=profile.rate_unit)
    if field_name in _DATE_FIELDS:
        return to_date(raw)
    if field_name in _INT_FIELDS:
        return to_int(raw)
    return to_text(raw)


def _row_is_data(row: Sequence[Any], cols: dict[str, int], spec: SheetSpec) -> bool:
    for guard in spec.row_guard:
        if to_text(_cell(row, cols.get(guard))) in (None, ""):
            return False
    for guard in spec.row_guard_text:
        if not isinstance(_cell(row, cols.get(guard)), str):
            return False
    return True


def _read_sheet(
    wb: WorkbookSource, spec: SheetSpec, profile: PortfolioProfile
) -> tuple[list[dict], int, list[ReadIssue]]:
    """Generic tabular read -> list of coerced field dicts (+ ``_row`` number)."""
    issues: list[ReadIssue] = []
    rows = wb.rows(spec.sheet)
    if not rows:
        raise ProfileMismatch(f"sheet {spec.sheet!r} is missing or empty")
    if len(rows) < spec.header_row:
        raise ProfileMismatch(
            f"sheet {spec.sheet!r} has {len(rows)} rows, header row {spec.header_row} not reachable"
        )
    cols = resolve_columns(rows[spec.header_row - 1], spec)
    out: list[dict] = []
    skipped = 0
    for offset, row in enumerate(rows[spec.first_data_row - 1:]):
        row_number = spec.first_data_row + offset
        if not any(c is not None and str(c).strip() != "" for c in row):
            continue  # entirely blank
        if not _row_is_data(row, cols, spec):
            skipped += 1
            continue
        record: dict[str, Any] = {"_row": row_number}
        for canonical, idx in cols.items():
            record[canonical] = _coerce(canonical, _cell(row, idx), profile)
        out.append(record)
    return out, skipped, issues


# ---------------------------------------------------------------------------
# Entity readers
# ---------------------------------------------------------------------------


def read_accounts(
    wb: WorkbookSource, profile: PortfolioProfile
) -> tuple[list[PortfolioAccount], int, list[ReadIssue]]:
    assert profile.accounts is not None
    records, skipped, issues = _read_sheet(wb, profile.accounts, profile)
    accounts: list[PortfolioAccount] = []
    for r in records:
        accounts.append(
            PortfolioAccount(
                row_number=r["_row"],
                account_number=str(r.get("account_number")),
                vendor_code=r.get("vendor_code"),
                provider_name=r.get("provider_name"),
                source_status=(r.get("status") or "").strip().upper(),
                source_sub_status=(r.get("sub_status") or "").strip().upper(),
                borrower_name=r.get("borrower_name"),
                co_borrower_name=r.get("co_borrower_name"),
                days_past_due=r.get("days_past_due") or 0,
                sales_value_cents=r.get("sales_value"),
                insurance_cents=r.get("insurance"),
                down_payment_cents=r.get("down_payment"),
                amount_financed_cents=r.get("amount_financed"),
                term_months=r.get("term_months"),
                annual_rate_bps=r.get("annual_rate"),
                regular_payment_cents=r.get("regular_payment"),
                payment_frequency=r.get("payment_frequency"),
                cost_of_borrowing_cents=r.get("cost_of_borrowing"),
                first_payment_date=r.get("first_payment_date"),
                final_payment_date=r.get("final_payment_date"),
                days_in_year=r.get("days_in_year"),
                origination_date=r.get("origination_date"),
                origination_type=r.get("origination_type"),
                payment_amount_total_cents=r.get("payment_amount_total"),
                fees_paid_cents=r.get("fees_paid"),
                interest_paid_cents=r.get("interest_paid"),
                principal_paid_cents=r.get("principal_paid"),
                fees_balance_cents=r.get("fees_balance"),
                interest_balance_cents=r.get("interest_balance"),
                principal_balance_cents=r.get("principal_balance"),
                total_owed_cents=r.get("total_owed"),
                next_due_date=r.get("next_due_date"),
                last_transaction_date=r.get("last_transaction_date"),
                last_transaction_type=r.get("last_transaction_type"),
                close_date=r.get("close_date"),
                close_type=r.get("close_type"),
                nsf_return_count=r.get("nsf_return_count") or 0,
                deferment_count=r.get("deferment_count") or 0,
            )
        )
    return accounts, skipped, issues


def read_transactions(
    wb: WorkbookSource, profile: PortfolioProfile
) -> tuple[list[PortfolioTransaction], int, list[ReadIssue]]:
    if profile.transactions is None:
        return [], 0, []
    records, skipped, issues = _read_sheet(wb, profile.transactions, profile)
    txns: list[PortfolioTransaction] = []
    for r in records:
        txn = PortfolioTransaction(
            row_number=r["_row"],
            account_number=str(r.get("account_number")),
            source_type=(r.get("type") or "").strip().upper(),
            date=r.get("date"),
            transaction_number=r.get("transaction_number"),
            vendor_code=r.get("vendor_code"),
            provider_name=r.get("provider_name"),
            payment_cents=r.get("payment"),
            fees_charged_cents=r.get("fees_charged"),
            fees_paid_cents=r.get("fees_paid"),
            fees_balance_cents=r.get("fees_balance"),
            accrued_interest_cents=r.get("accrued_interest"),
            interest_due_cents=r.get("interest_due"),
            interest_paid_cents=r.get("interest_paid"),
            interest_balance_cents=r.get("interest_balance"),
            principal_paid_cents=r.get("principal_paid"),
            principal_balance_cents=r.get("principal_balance"),
            total_owed_cents=r.get("total_owed"),
            comment=r.get("comment"),
        )
        if txn.date is None:
            issues.append(
                ReadIssue(profile.transactions.sheet, txn.row_number,
                          "transaction has no parseable date — not imported")
            )
            continue
        txns.append(txn)
    return txns, skipped, issues


def read_vendors(
    wb: WorkbookSource, profile: PortfolioProfile
) -> tuple[list[PortfolioVendor], list[ReadIssue]]:
    """Read a TRANSPOSED vendor sheet (labels down a column, one vendor per column)."""
    spec: Optional[TransposedSheetSpec] = profile.vendors
    if spec is None:
        return [], []
    issues: list[ReadIssue] = []
    rows = wb.rows(spec.sheet)
    if not rows:
        return [], [ReadIssue(spec.sheet, None, "vendor sheet missing or empty")]

    # label text -> 0-based row index
    label_rows: dict[str, int] = {}
    for i, row in enumerate(rows):
        label = _normalize_header(_cell(row, spec.label_column))
        if label and label not in label_rows:
            label_rows[label] = i

    guard_row = label_rows.get(_normalize_header(spec.labels.get(spec.entity_guard, "")))
    if guard_row is None:
        return [], [ReadIssue(spec.sheet, None,
                              f"vendor sheet has no {spec.entity_guard!r} label row")]

    width = max((len(r) for r in rows), default=0)
    vendors: list[PortfolioVendor] = []
    for col in range(spec.first_entity_column, width):
        if to_text(_cell(rows[guard_row], col)) is None:
            continue
        v = PortfolioVendor()
        for canonical, label in spec.labels.items():
            ri = label_rows.get(_normalize_header(label))
            if ri is None:
                continue
            raw = _cell(rows[ri], col)
            setattr(v, canonical, to_date(raw) if canonical in _DATE_FIELDS else to_text(raw))
        for canonical, (label, span) in spec.list_blocks.items():
            ri = label_rows.get(_normalize_header(label))
            if ri is None:
                continue
            values: list[str] = []
            for r in range(ri, min(ri + span, len(rows))):
                # A blank label cell means the block continues; a NEW label ends it.
                if r != ri and _normalize_header(_cell(rows[r], spec.label_column)):
                    break
                item = to_text(_cell(rows[r], col))
                if item:
                    values.append(item)
            setattr(v, canonical, values)
        vendors.append(v)
    return vendors, issues


def read_workbook(wb: WorkbookSource, profile: PortfolioProfile) -> WorkbookReadResult:
    """Read every entity the profile describes. Raises ``ProfileMismatch`` when
    the profile plainly does not fit the file (missing sheet / required column)."""
    problems = profile.validate()
    if problems:
        raise ProfileMismatch("; ".join(problems))
    result = WorkbookReadResult(profile_name=profile.name)
    accounts, acc_skipped, acc_issues = read_accounts(wb, profile)
    txns, txn_skipped, txn_issues = read_transactions(wb, profile)
    vendors, vendor_issues = read_vendors(wb, profile)
    result.accounts = accounts
    result.transactions = txns
    result.vendors = vendors
    result.skipped_rows = acc_skipped + txn_skipped
    result.issues = [*acc_issues, *txn_issues, *vendor_issues]
    return result


def read_file(path_or_bytes: Any, profile: PortfolioProfile) -> WorkbookReadResult:
    wb = ExcelWorkbook(path_or_bytes)
    try:
        return read_workbook(wb, profile)
    finally:
        wb.close()


def split_person_name(raw: Optional[str], name_format: str) -> tuple[Optional[str], Optional[str]]:
    """"``Last, First``" -> ``("First", "Last")``, per the profile's declaration."""
    text = (raw or "").strip()
    if not text:
        return None, None
    if name_format == "last_comma_first" and "," in text:
        last, _, first = text.partition(",")
        return (first.strip() or None), (last.strip() or None)
    if name_format == "first_last":
        parts = text.split()
        if len(parts) >= 2:
            return parts[0], " ".join(parts[1:])
        return text, None
    parts = [p.strip() for p in text.split(",") if p.strip()]
    if len(parts) >= 2:
        return parts[1], parts[0]
    words = text.split()
    if len(words) >= 2:
        return " ".join(words[:-1]), words[-1]
    return text, None


def iter_by_account(
    transactions: Iterable[PortfolioTransaction],
) -> dict[str, list[PortfolioTransaction]]:
    """Group transactions by account, preserving the file's row order."""
    out: dict[str, list[PortfolioTransaction]] = {}
    for t in transactions:
        out.setdefault(t.account_number, []).append(t)
    return out
